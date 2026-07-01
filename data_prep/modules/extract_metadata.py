"""
Professional Metadata Extraction Module for RAG Pipelines
--------------------------------------------------------
- Extracts rich metadata from PDFs, DOCX, PPTX, XLSX, and more.
- Adds source_link, content_hash, extraction_time, and counts (pages, slides, etc.).
- Logs a summary at the end.
"""
import json
import mimetypes
import hashlib
import logging
from pathlib import Path
from typing import Dict, Any
from datetime import datetime
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

OUTPUT_DIR = Path("./data_prep/processed/metadata")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s')
logger = logging.getLogger("extract_metadata")

def compute_file_hash(path: Path, block_size: int = 65536) -> str:
    sha = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(block_size), b''):
            sha.update(block)
    return sha.hexdigest()

def build_source_link(path: Path) -> str:
    return f"file:///{path.resolve().as_posix()}"

def get_pdf_metadata(path: Path) -> Dict[str, Any]:
    try:
        reader = PdfReader(str(path))
        info = reader.metadata or {}
        return {
            "title": info.get("/Title", ""),
            "author": info.get("/Author", ""),
            "subject": info.get("/Subject", ""),
            "producer": info.get("/Producer", ""),
            "num_pages": len(reader.pages)
        }
    except Exception as e:
        return {"error": str(e)}

def get_docx_metadata(path: Path) -> Dict[str, Any]:
    try:
        doc = Document(path)
        cp = doc.core_properties
        return {
            "title": cp.title,
            "author": cp.author,
            "created": cp.created.isoformat() if cp.created else "",
            "num_sections": len(doc.sections)
        }
    except Exception as e:
        return {"error": str(e)}

def get_pptx_metadata(path: Path) -> Dict[str, Any]:
    try:
        ppt = Presentation(path)
        cp = ppt.core_properties
        return {
            "title": cp.title,
            "author": cp.author,
            "created": cp.created.isoformat() if cp.created else "",
            "num_slides": len(ppt.slides)
        }
    except Exception as e:
        return {"error": str(e)}

def get_xlsx_metadata(path: Path) -> Dict[str, Any]:
    try:
        wb = load_workbook(path, read_only=True)
        props = wb.properties
        return {
            "title": props.title,
            "author": props.creator,
            "created": props.created.isoformat() if props.created else "",
            "num_sheets": len(wb.worksheets)
        }
    except Exception as e:
        return {"error": str(e)}

def generic_metadata(path: Path) -> Dict[str, Any]:
    return {
        "file_name": path.name,
        "file_path": str(path.resolve()),
        "file_size": path.stat().st_size,
        "file_type": mimetypes.guess_type(str(path))[0] or path.suffix,
        "modified": datetime.fromtimestamp(path.stat().st_mtime).isoformat(),
        "content_hash": compute_file_hash(path),
        "extraction_time": datetime.now().isoformat(),
        "source_link": build_source_link(path)
    }

EXTRACTOR_MAP = {
    ".pdf": get_pdf_metadata,
    ".docx": get_docx_metadata,
    ".pptx": get_pptx_metadata,
    ".xlsx": get_xlsx_metadata,
}

def extract_metadata_from_all(raw_dir: Path = Path("raw_files"), output_dir: Path = OUTPUT_DIR):
    output_dir.mkdir(parents=True, exist_ok=True)
    summary = {"files": 0, "errors": 0}
    for file_path in raw_dir.rglob("*"):
        if not file_path.is_file():
            continue
        ext = file_path.suffix.lower()
        extractor = EXTRACTOR_MAP.get(ext, lambda _: {})
        meta = generic_metadata(file_path)
        meta.update(extractor(file_path))
        output_path = output_dir / (file_path.stem + ".json")
        try:
            with open(output_path, "w", encoding="utf-8") as f:
                json.dump(meta, f, indent=2, ensure_ascii=False)
            logger.info(f"[meta] Extracted: {file_path.name}")
            summary["files"] += 1
        except Exception as e:
            logger.error(f"[meta] Failed to write {output_path}: {e}")
            summary["errors"] += 1
    logger.info(f"[summary] Files processed: {summary['files']}, Errors: {summary['errors']}")

if __name__ == "__main__":
    extract_metadata_from_all()
